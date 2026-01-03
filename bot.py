import os
import time
import logging
from contextlib import closing
from datetime import datetime, timezone

import psycopg2
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update, InputFile
from telegram.ext import (
    Application, CommandHandler, CallbackQueryHandler, ContextTypes,
    ConversationHandler, MessageHandler, filters
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

TOKEN = "8367905898:AAEimA-3iLi-JqP9r4cJPnOzYE-L4eYsk-U"
ADMIN_CHAT_ID = 5443870760

INVITE_LINK = "https://t.me/NaturalSense"
CONTACT_EMAIL = "naturalsense.pr@gmail.com"
CONTACT_TG = "@NScollab"
COVER_PATH = "cover.jpg"

TITLE = "✨ NS · Natural Sense"
SUBTITLE = "Распаковки · Бренды · Обзоры\nСравнения · Новости"
STATUS = "🔒 Закрытый доступ"
MAIN_CAPTION = f"{TITLE}\n\n{SUBTITLE}\n\n{STATUS}"
SUCCESS_CAPTION = f"{TITLE}\n\n✅ Заявка отправлена.\nМы свяжемся с вами по указанному контакту.\n\n{STATUS}"

ABOUT_TEXT = "ℹ️ *О канале*\n\nNatural Sense — обзоры косметики, брендов и новинок."
COLLAB_TEXT = "🤝 *Сотрудничество*\n\nВыберите удобный способ связи или оставьте заявку:"
EXIT_TEXT = "❌ Вы вышли из меню.\n\nНажми /start чтобы открыть меню снова."

MIN_SECONDS_BETWEEN_FORM_MSG = 0.7
COOLDOWN_SECONDS = 15 * 60
MAX_INPUT_LEN = 500

FORM_INTRO = (
    "📦 *Заявка на интеграцию*\n\n"
    "Отвечайте по шагам. Чтобы отменить — отправьте /cancel."
)
Q = [
    "1) 🏷 Название бренда / компании?",
    "2) 📦 Что продвигаем? (продукт/линейка/ссылка)",
    "3) 💰 Бюджет / условия? (сумма, бартер, % и т.д.)",
    "4) 📞 Контакт для связи (email или Telegram @username)",
    "5) 📝 Доп. детали (необязательно). Если нет — напиши: -",
]
K = ["brand", "product", "budget", "contact", "extra"]
S0, S1, S2, S3, S4 = range(5)

def _now(): return time.time()

def _flood_ok(ctx, uid):
    m = ctx.application.bot_data.setdefault("antiflood", {})
    t = _now()
    if (t - m.get(uid, 0)) < MIN_SECONDS_BETWEEN_FORM_MSG:
        return False
    m[uid] = t
    return True

def _cooldown_left(ctx, uid):
    m = ctx.application.bot_data.setdefault("cooldowns", {})
    left = int(m.get(uid, 0) - _now())
    return left if left > 0 else 0

def _set_cooldown(ctx, uid):
    ctx.application.bot_data.setdefault("cooldowns", {})[uid] = _now() + COOLDOWN_SECONDS

def _clean(t: str) -> str:
    t = (t or "").strip()
    return t[:MAX_INPUT_LEN]

def _gmail_compose(email: str) -> str:
    return f"https://mail.google.com/mail/?view=cm&to={email}"

def _tg_link(username: str) -> str:
    return f"https://t.me/{username.replace('@','').strip()}"

def kb_main():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🔐 Войти в канал", url=INVITE_LINK)],
        [InlineKeyboardButton("ℹ️ О канале", callback_data="about")],
        [InlineKeyboardButton("🤝 Сотрудничество", callback_data="collab")],
        [InlineKeyboardButton("❌ Выйти", callback_data="exit")],
    ])

def kb_back():
    return InlineKeyboardMarkup([[InlineKeyboardButton("⬅ Назад", callback_data="back")]])

def kb_collab():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📧 Написать на Email", url=_gmail_compose(CONTACT_EMAIL))],
        [InlineKeyboardButton("💬 Написать в Telegram", url=_tg_link(CONTACT_TG))],
        [InlineKeyboardButton("📦 Предложить интеграцию", callback_data="integration")],
        [InlineKeyboardButton("⬅ Назад", callback_data="back")],
    ])

async def send_menu_photo(ctx: ContextTypes.DEFAULT_TYPE, chat_id: int, caption: str):
    if not os.path.exists(COVER_PATH):
        await ctx.bot.send_message(chat_id=chat_id, text="❗ cover.jpg не найден рядом с bot.py")
        return
    with open(COVER_PATH, "rb") as f:
        await ctx.bot.send_photo(chat_id=chat_id, photo=InputFile(f), caption=caption, reply_markup=kb_main())

def db_conn():
    url = os.getenv("DATABASE_URL")
    if not url:
        raise RuntimeError("DATABASE_URL not set")
    return psycopg2.connect(url, sslmode="require")

def db_init():
    with closing(db_conn()) as conn:
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS integrations (
                    id BIGSERIAL PRIMARY KEY,
                    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    user_id BIGINT NOT NULL,
                    username TEXT,
                    brand TEXT,
                    product TEXT,
                    budget TEXT,
                    contact TEXT,
                    extra TEXT
                );
            """)
        conn.commit()

def db_insert(rec: dict):
    with closing(db_conn()) as conn:
        with conn.cursor() as cur:
            cur.execute("""
                INSERT INTO integrations (user_id, username, brand, product, budget, contact, extra)
                VALUES (%s,%s,%s,%s,%s,%s,%s)
            """, (
                rec["user_id"], rec.get("username"),
                rec.get("brand"), rec.get("product"), rec.get("budget"),
                rec.get("contact"), rec.get("extra"),
            ))
        conn.commit()

def db_fetch_all():
    with closing(db_conn()) as conn:
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, created_at, user_id, username, brand, product, budget, contact, extra
                FROM integrations
                ORDER BY created_at DESC
            """)
            return cur.fetchall()

def build_xlsx(rows, path="integrations.xlsx"):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "integrations"
    headers = ["id","created_at","user_id","username","brand","product","budget","contact","extra"]
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for i in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    wb.save(path)

async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await send_menu_photo(ctx, update.effective_chat.id, MAIN_CAPTION)

async def menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    if q.data == "exit":
        await q.message.edit_caption(caption=EXIT_TEXT, reply_markup=None)
    elif q.data == "back":
        await q.message.edit_caption(caption=MAIN_CAPTION, reply_markup=kb_main())
    elif q.data == "about":
        await q.message.edit_caption(caption=ABOUT_TEXT, parse_mode="Markdown", reply_markup=kb_back())
    elif q.data == "collab":
        await q.message.edit_caption(caption=COLLAB_TEXT, parse_mode="Markdown", reply_markup=kb_collab())

async def form_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    left = _cooldown_left(ctx, uid)
    if left > 0:
        mins = (left + 59) // 60
        await q.message.reply_text(f"⏳ Подождите {mins} мин. перед повторной отправкой заявки.")
        return ConversationHandler.END
    ctx.user_data["form"] = {}
    ctx.user_data["step"] = 0
    await q.message.reply_text(FORM_INTRO + "\n\n" + Q[0], parse_mode="Markdown")
    return S0

async def form_step(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    if not _flood_ok(ctx, uid):
        return int(ctx.user_data.get("step", 0))

    step = int(ctx.user_data.get("step", 0))
    text = _clean(update.message.text)

    if step == 4 and text == "-":
        text = "—"

    ctx.user_data["form"][K[step]] = text
    step += 1
    ctx.user_data["step"] = step

    if step < 5:
        await update.message.reply_text(Q[step])
        return step

    user = update.effective_user
    form = ctx.user_data.get("form", {})
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    rec = {
        "user_id": user.id,
        "username": user.username,
        **form,
    }

    try:
        db_insert(rec)
    except Exception as e:
        logging.exception("DB insert failed: %s", e)
        try:
            await ctx.bot.send_message(chat_id=ADMIN_CHAT_ID, text=f"🚨 DB insert failed:\n{e}")
        except Exception:
            pass

    admin_text = (
        "📦 Новая заявка на интеграцию\n\n"
        f"🕒 {ts}\n"
        f"👤 User: @{user.username if user.username else '—'} (ID: {user.id})\n\n"
        f"🏷 Бренд: {rec.get('brand','—')}\n"
        f"📦 Продукт: {rec.get('product','—')}\n"
        f"💰 Условия: {rec.get('budget','—')}\n"
        f"📞 Контакт: {rec.get('contact','—')}\n"
        f"📝 Детали: {rec.get('extra','—')}\n"
    )
    try:
        await ctx.bot.send_message(chat_id=ADMIN_CHAT_ID, text=admin_text)
    except Exception:
        logging.exception("Admin notify failed")

    _set_cooldown(ctx, uid)
    await send_menu_photo(ctx, update.effective_chat.id, SUCCESS_CAPTION)

    ctx.user_data.pop("form", None)
    ctx.user_data.pop("step", None)
    return ConversationHandler.END

async def cancel(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    ctx.user_data.pop("form", None)
    ctx.user_data.pop("step", None)
    await update.message.reply_text("❌ Заявка отменена. Нажми /start чтобы вернуться в меню.")
    return ConversationHandler.END

async def export_excel(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_CHAT_ID:
        return
    try:
        rows = db_fetch_all()
        path = "integrations.xlsx"
        build_xlsx(rows, path)
        with open(path, "rb") as f:
            await ctx.bot.send_document(
                chat_id=ADMIN_CHAT_ID,
                document=f,
                filename="integrations.xlsx",
                caption=f"📊 Экспорт заявок: {len(rows)} шт.",
            )
    except Exception as e:
        logging.exception("Export failed: %s", e)
        try:
            await ctx.bot.send_message(chat_id=ADMIN_CHAT_ID, text=f"🚨 Export failed:\n{e}")
        except Exception:
            pass

async def on_startup(app: Application):
    try:
        db_init()
    except Exception as e:
        logging.exception("DB init failed: %s", e)
        try:
            await app.bot.send_message(chat_id=ADMIN_CHAT_ID, text=f"🚨 DB init failed:\n{e}")
        except Exception:
            pass
    try:
        await app.bot.send_message(chat_id=ADMIN_CHAT_ID, text="✅ BOT STARTED")
    except Exception:
        pass

async def on_shutdown(app: Application):
    try:
        await app.bot.send_message(chat_id=ADMIN_CHAT_ID, text="🛑 BOT STOPPED / RESTARTING")
    except Exception:
        pass

async def errors(update: object, ctx: ContextTypes.DEFAULT_TYPE):
    err = str(ctx.error)
    if "terminated by other getUpdates request" in err:
        logging.warning("Conflict: another getUpdates is running.")
        return
    logging.exception("Unhandled error:", exc_info=ctx.error)
    try:
        await ctx.bot.send_message(chat_id=ADMIN_CHAT_ID, text=f"🚨 Bot error:\n{err}")
    except Exception:
        pass

def main():
    app = Application.builder().token(TOKEN).build()
    app.post_init = on_startup
    app.post_shutdown = on_shutdown

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("export_excel", export_excel))

    app.add_handler(CallbackQueryHandler(menu, pattern=r"^(back|about|collab|exit)$"))

    conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(form_start, pattern=r"^integration$")],
        states={
            S0: [MessageHandler(filters.TEXT & ~filters.COMMAND, form_step)],
            S1: [MessageHandler(filters.TEXT & ~filters.COMMAND, form_step)],
            S2: [MessageHandler(filters.TEXT & ~filters.COMMAND, form_step)],
            S3: [MessageHandler(filters.TEXT & ~filters.COMMAND, form_step)],
            S4: [MessageHandler(filters.TEXT & ~filters.COMMAND, form_step)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )
    app.add_handler(conv)

    app.add_error_handler(errors)
    app.run_polling()

if __name__ == "__main__":
    main()
